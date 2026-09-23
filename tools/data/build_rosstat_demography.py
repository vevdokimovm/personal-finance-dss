"""Демография: браки, разводы, возрастной состав → `data/`.

Оригиналы (`~/raw-originals/finpilot-data/rosstat_demo/`): `demo31_2023.xlsx`
(браки), `demo32_2023.xlsx` (разводы), `demo14.xlsx` (распределение
населения по возрастным группам).

Ядру это нужно как частота жизненных событий, вокруг которых строятся
цели (свадьба, рождение ребёнка, развод как шок), и как возрастной
профиль пользователя.

Запуск: `python -m tools.data.build_rosstat_demography`
"""
from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path

import openpyxl

from tools.data.common import Row, raw_root, to_float, write_series

logger = logging.getLogger(__name__)

SECTIONS = {
    "Всё население": "total",
    "Городское население": "urban",
    "Сельское население": "rural",
}


def read_events(path: Path, event: str) -> list[Row]:
    """Прочитать таблицу «годы → число событий и коэффициент».

    Args:
        path: путь к книге `demo31`/`demo32`.
        event: `marriages` или `divorces`.

    Returns:
        Строки ряда: абсолютное число и коэффициент на 1000 человек.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[Row] = []
    for raw in book[book.sheetnames[0]].iter_rows(values_only=True):
        label = str(raw[0] or "").strip()
        if not re.fullmatch(r"\d{4}", label):
            continue
        total = to_float(raw[1]) if len(raw) > 1 else None
        per_1000 = to_float(raw[2]) if len(raw) > 2 else None
        date = f"{label}-12-31"
        if total is not None:
            out.append(Row(date, total, "events", "RU", f"{event}|total"))
        if per_1000 is not None:
            out.append(Row(date, per_1000, "per_1000_population", "RU",
                           f"{event}|rate"))
    return out


def read_age_groups(path: Path) -> list[Row]:
    """Прочитать распределение населения по возрастным группам.

    Args:
        path: путь к `demo14.xlsx`.

    🔴 Таблица состоит из трёх одинаковых блоков: всё население,
    городское и сельское. Разрез блока ставится первым уровнем
    `breakdown` (`total|0 – 4`, `urban|0 – 4`, `rural|0 – 4`); без него
    три блока сливались под одним ключом (поймано прогоном B2).

    Returns:
        Строки ряда: численность группы в тысячах человек на 1 января.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[book.sheetnames[0]].iter_rows(values_only=True))
    header_index = next(i for i, raw in enumerate(rows)
                        if sum(1 for cell in raw
                               if re.fullmatch(r"\d{4}", str(cell or "").strip()))
                        >= 3)
    years = {index: str(cell).strip()
             for index, cell in enumerate(rows[header_index])
             if re.fullmatch(r"\d{4}", str(cell or "").strip())}
    out: list[Row] = []
    section = "total"
    for raw in rows[header_index + 1:]:
        label = re.sub(r"\s+", " ", str(raw[0] or "").replace("\xa0", " ")).strip()
        if not label or label.lower().startswith(("в том числе", "из общей")):
            continue
        if label in SECTIONS:
            section, label = SECTIONS[label], "all"
        label = re.sub(r"\s*\d\)$", "", label)
        label = label[:1].lower() + label[1:]
        for index, year in years.items():
            value = to_float(raw[index]) if index < len(raw) else None
            if value is None:
                continue
            out.append(Row(f"{year}-01-01", value, "thousand_people", "RU",
                           f"{section}|{label}"))
    return out


def build(raw_dir: Path, out_dir: Path | None = None) -> None:
    """Собрать демографические ряды.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения.
    """
    folder = raw_dir / "rosstat_demo"
    events = read_events(folder / "demo31_2023.xlsx", "marriages")
    events += read_events(folder / "demo32_2023.xlsx", "divorces")
    events.sort(key=lambda item: (item.date, item.breakdown))
    write_series("rosstat_marriages_divorces_year.csv", events, out_dir)
    write_series("rosstat_population_age_groups.csv",
                 read_age_groups(folder / "demo14.xlsx"), out_dir)


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=raw_root())
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.raw_dir, args.out_dir)


if __name__ == "__main__":
    main()
