"""Обследование бюджетов домашних хозяйств (ОБДХ) → `data/`.

Оригиналы — архивы `.rar` в `~/raw-originals/finpilot-data/rosstat_obdh/`
(итоги Росстата «Доходы, расходы и потребление домашних хозяйств»).
Распаковка выполняется во временный каталог через `unar`; в репозиторий
кладутся только два ряда:

  * `rosstat_obdh_decile_level_year.csv` — рубли в месяц на члена
    домохозяйства по децильным группам (лист `1.1.4`);
  * `rosstat_obdh_decile_structure_year.csv` — структура
    потребительских расходов в процентах (лист `1.5.6`).

Берутся только годовые выпуски (каталог вида `DRP_2025`): квартальные
и накопленные выпуски (`DRP_2026_1`) дают в тех же колонках период,
а не год, и при слиянии портят годовой ряд.

Запуск: `python -m tools.data.build_rosstat_obdh`
"""
from __future__ import annotations

import argparse
import logging
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl

from tools.data.common import Row, raw_root, to_float, write_series

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

DECILES = ["Первая", "Вторая", "Третья", "Четвертая", "Пятая", "Шестая",
           "Седьмая", "Восьмая", "Девятая", "Десятая"]
DECILE_CODES = {name: f"decile_{index}" for index, name in enumerate(DECILES, 1)}
SKIP_PREFIXES = ("из них", "в том числе", "справочно")


def unpack(archive: Path, target: Path) -> None:
    """Распаковать архив ОБДХ, включая вложенные архивы.

    Args:
        archive: путь к `.rar`.
        target: каталог назначения.

    Raises:
        RuntimeError: если `unar` не установлен.
    """
    if shutil.which("unar") is None:
        raise RuntimeError("нужен распаковщик unar: brew install unar")
    subprocess.run(["unar", "-q", "-o", str(target), str(archive)],
                   check=True, capture_output=True)
    for nested in sorted(target.rglob("*.rar")):
        subprocess.run(["unar", "-q", "-o", str(nested.parent), str(nested)],
                       check=True, capture_output=True)
        nested.unlink()


def is_vague_label(label: str) -> bool:
    """Проверить, что подпись строки сама себя не называет.

    В таблицах Росстата часть строк подписана «в том числе», «из них»,
    «справочно»: что именно там измеряется, видно только из строки выше.

    Args:
        label: подпись строки.

    Returns:
        True, если подпись неинформативна сама по себе.
    """
    return label.lower().startswith(SKIP_PREFIXES)


def read_decile_sheet(sheet: Worksheet, unit: str) -> list[Row]:
    """Разобрать лист с колонками «дециль × год».

    Args:
        sheet: лист книги ОБДХ (`1.1.4` или `1.5.6`).
        unit: единица измерения значений.

    Returns:
        Строки ряда с разрезом «дециль | показатель».
    """
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (i for i, raw in enumerate(rows)
         if raw and any(str(cell).strip() in DECILES
                        for cell in raw if cell)), None)
    if header_index is None:
        return []
    deciles: dict[int, str] = {}
    current = ""
    for index, cell in enumerate(rows[header_index]):
        name = str(cell).strip() if cell else ""
        if name in DECILES:
            current = name
        if current:
            deciles[index] = current
    years_row = rows[header_index + 1]
    years: dict[int, str] = {}
    for index, cell in enumerate(years_row):
        match = re.search(r"(\d{4})", str(cell or ""))
        if match and index in deciles:
            years[index] = match.group(1)
    out: list[Row] = []
    parent = ""
    for raw in rows[header_index + 2:]:
        label = re.sub(r"\s+", " ", str(raw[0] or "")).strip()
        if not label:
            continue
        if is_vague_label(label):
            # Строка есть, а имени у неё нет: сохраняем с указанием, к чему
            # она уточнение, чтобы число не потерялось и не соврало.
            label = f"{label} (уточнение к: {parent})" if parent else label
        else:
            parent = label
        for index, year in years.items():
            value = to_float(raw[index]) if index < len(raw) else None
            if value is None:
                continue
            out.append(Row(f"{year}-12-31", value, unit, "RU",
                           f"{DECILE_CODES[deciles[index]]}|{label}"))
    return out


def build(raw_dir: Path, out_dir: Path | None = None) -> None:
    """Распаковать архивы ОБДХ и собрать два ряда.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения.
    """
    archives = sorted((raw_dir / "rosstat_obdh").glob("*.rar"))
    level: dict[tuple[str, str], Row] = {}
    structure: dict[tuple[str, str], Row] = {}
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        for archive in archives:
            unpack(archive, work / archive.stem)
        for book_path in sorted(work.rglob("*РАЗДЕЛ_1.xlsx")):
            if not re.fullmatch(r"DRP_\d{4}", book_path.parent.name):
                logger.info("пропускаю квартальный выпуск %s", book_path.parent.name)
                continue
            book = openpyxl.load_workbook(book_path, read_only=True,
                                          data_only=True)
            if "1.1.4" in book.sheetnames:
                for row in read_decile_sheet(book["1.1.4"], "rub_per_member_month"):
                    level[(row.date, row.breakdown)] = row
            if "1.5.6" in book.sheetnames:
                for row in read_decile_sheet(book["1.5.6"], "percent_of_spending"):
                    structure[(row.date, row.breakdown)] = row
    write_series("rosstat_obdh_decile_level_year.csv",
                 [level[key] for key in sorted(level)], out_dir)
    write_series("rosstat_obdh_decile_structure_year.csv",
                 [structure[key] for key in sorted(structure)], out_dir)


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=raw_root())
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.raw_dir, args.out_dir)


if __name__ == "__main__":
    main()
