"""Распределение населения по среднедушевому доходу → `data/`.

Источник: Росстат, «Неравенство и бедность», `NB_RD_1-2-8.xlsx`
(листы «субъекты РФ 1995» … «субъекты РФ 2025»).

Даёт два ряда:
  * `rosstat_income_distribution_year.csv` — доля населения в интервале
    дохода, по годам и субъектам;
  * `rosstat_income_lognormal_year.csv` — параметры логнормали,
    подогнанной по интервалам того же года (σ, медиана, квантили).

Запуск: `python -m tools.data.build_rosstat_income`
"""
from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl

from tools.data.common import Row, fit_lognormal, raw_root, to_float, write_series

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

SOURCE = "rosstat_ineq/NB_RD_1-2-8.xlsx"
QUANTILES = (0.1, 0.25, 0.5, 0.75, 0.9, 0.95)


def parse_bound(label: str) -> float | None:
    """Вытащить верхнюю границу интервала дохода из заголовка колонки.

    Args:
        label: заголовок вида `«10 000,1-14 000,0»` или `«свыше 100 000,0»`.

    Returns:
        Верхняя граница в рублях либо None для открытого сверху интервала.
    """
    text = label.replace("\xa0", " ").strip()
    if text.lower().startswith("свыше"):
        return None
    numbers = re.findall(r"\d[\d  ]*(?:,\d+)?", text)
    if not numbers:
        return None
    return to_float(numbers[-1])


def read_sheet(sheet: Worksheet) -> tuple[
        list[str], list[float | None], list[tuple[str, list[float | None]]]]:
    """Прочитать лист одного года.

    Args:
        sheet: лист книги Росстата.

    Returns:
        Кортеж из подписей интервалов, их верхних границ и списка пар
        «регион → доли по интервалам».
    """
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[3]]
    columns = [i for i, name in enumerate(header) if i >= 2 and name]
    labels = [header[i].replace("\xa0", " ") for i in columns]
    bounds = [parse_bound(name) for name in labels]
    regions: list[tuple[str, list[float | None]]] = []
    for raw in rows[4:]:
        name = str(raw[0]).strip() if raw and raw[0] is not None else ""
        if not name:
            continue
        shares = [to_float(raw[i]) if i < len(raw) else None for i in columns]
        if all(share is None for share in shares):
            continue
        regions.append((name, shares))
    return labels, bounds, regions


def build(source: Path, out_dir: Path | None = None) -> None:
    """Собрать оба ряда из книги Росстата.

    Args:
        source: путь к `NB_RD_1-2-8.xlsx`.
        out_dir: каталог назначения (по умолчанию `data/`).
    """
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    dist_rows: list[Row] = []
    fit_rows: list[Row] = []
    for sheet_name in book.sheetnames:
        match = re.search(r"(\d{4})", sheet_name)
        if not match or "субъект" not in sheet_name.lower():
            continue
        year = int(match.group(1))
        date = f"{year}-12-31"
        labels, bounds, regions = read_sheet(book[sheet_name])
        for region, shares in regions:
            for label, share in zip(labels, shares):
                if share is None or label.lower() == "всего":
                    continue
                dist_rows.append(
                    Row(date, share, "percent_of_population", region, label)
                )
            fit_bounds = [b for b in bounds if b is not None]
            fit_shares = [0.0 if s is None else s
                          for s, lb in zip(shares, labels)
                          if lb.lower() != "всего"]
            if len(fit_shares) != len(fit_bounds) + 1:
                continue
            try:
                fit = fit_lognormal(fit_bounds, fit_shares)
            except ValueError:
                logger.warning("%s %s: подгонка не сошлась", year, region)
                continue
            fit_rows.append(Row(date, fit.sigma, "sigma_of_log", region, "sigma"))
            fit_rows.append(Row(date, fit.mu, "mu_of_log", region, "mu"))
            fit_rows.append(Row(date, fit.median, "rub_per_month", region, "median"))
            fit_rows.append(Row(date, fit.mean, "rub_per_month", region, "mean"))
            fit_rows.append(Row(date, fit.r_squared, "r2", region, "r_squared"))
            for level in QUANTILES:
                fit_rows.append(
                    Row(date, fit.quantile(level), "rub_per_month", region,
                        f"q{int(level * 100)}")
                )
    write_series("rosstat_income_distribution_year.csv", dist_rows, out_dir)
    write_series("rosstat_income_lognormal_year.csv", fit_rows, out_dir)


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=raw_root() / SOURCE)
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.source, args.out_dir)


if __name__ == "__main__":
    main()
