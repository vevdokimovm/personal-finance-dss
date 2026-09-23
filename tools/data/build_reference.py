"""Справочные и контекстные ряды → `data/`.

Собирает то, что по отдельности слишком мелко для своего скрипта:

  * `world_benchmarks.csv` — мировой контекст (FRED и World Bank):
    нужен как проверка «наши пороги не абсурдны на фоне мира», а не как
    источник российских чисел;
  * `ofz_zcyc_params.csv` — параметры кривой бескупонной доходности ОФЗ
    (Нельсон—Сигель—Свенссон), по одному снимку на дату;
  * `cbr_mpl_limits.csv` — макропруденциальные лимиты Банка России
    (регуляторные пороги по ПДН как справочник рантайма);
  * `cbr_deposit_rate_top10_decade.csv` — максимальная ставка десяти
    банков с наибольшим объёмом вкладов, по декадам;
  * `lbma_gold_pm_daily.csv` — цена золота (LBMA PM fixing) с 1968 года:
    самый длинный доступный дневной ряд для стресс-теста долгих целей.

Запуск: `python -m tools.data.build_reference`
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import logging
import re
from pathlib import Path

import openpyxl

from tools.data.common import Row, raw_root, to_float, write_series

logger = logging.getLogger(__name__)

FRED_SERIES = {
    "TDSP": ("percent_of_income", "US", "household_debt_service_ratio"),
    "PSAVERT": ("percent_of_income", "US", "personal_saving_rate"),
    "SIPOVGINIUSA": ("gini_index", "US", "gini"),
    "MORTGAGE30US": ("percent_annual", "US", "mortgage_30y_rate"),
    "CPIAUCSL": ("index_1982_84", "US", "cpi"),
    "DGS10": ("percent_annual", "US", "treasury_10y"),
    "HDTGPDUSQ163N": ("percent_of_gdp", "US", "household_debt_to_gdp"),
    "MEHOINUSA672N": ("usd_2024_per_year", "US", "median_household_income"),
    "VIXCLS": ("index_points", "US", "vix"),
    "SP500": ("index_points", "US", "sp500"),
}
WORLD_BANK_SERIES = {
    "SI.POV.GINI": ("gini_index", "gini"),
    "NY.GNS.ICTR.ZS": ("percent_of_gdp", "gross_savings"),
    "FR.INR.LEND": ("percent_annual", "lending_rate"),
    "FR.INR.DPST": ("percent_annual", "deposit_rate"),
    "FP.CPI.TOTL.ZG": ("percent_annual", "inflation"),
    "SP.DYN.LE00.IN": ("years", "life_expectancy"),
    "NY.GDP.PCAP.CD": ("usd_per_capita", "gdp_per_capita"),
    "SL.UEM.TOTL.ZS": ("percent_of_labour", "unemployment"),
}
DECADE = {"I": "01", "II": "11", "III": "21"}


def read_world(raw_dir: Path) -> list[Row]:
    """Собрать мировой контекст из FRED и World Bank.

    Args:
        raw_dir: каталог оригиналов.

    Returns:
        Строки ряда с регионом-страной и названием показателя.
    """
    rows: list[Row] = []
    for code, (unit, region, name) in FRED_SERIES.items():
        path = raw_dir / "fred" / f"fred_{code}.csv"
        if not path.exists():
            logger.warning("нет файла %s", path)
            continue
        with path.open(encoding="utf-8") as handle:
            for record in csv.DictReader(handle):
                date = record.get("observation_date", "")
                value = to_float(record.get(code))
                if date and value is not None:
                    rows.append(Row(date, value, unit, region, f"fred|{name}"))
    for code, (unit, name) in WORLD_BANK_SERIES.items():
        path = raw_dir / "worldbank" / f"wb_{code}.json"
        if not path.exists():
            logger.warning("нет файла %s", path)
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        for record in payload[1]:
            value = record.get("value")
            if value is None:
                continue
            rows.append(Row(f"{record['date']}-12-31", float(value), unit,
                            record["countryiso3code"], f"worldbank|{name}"))
    rows.sort(key=lambda item: (item.breakdown, item.region, item.date))
    return rows


def read_zcyc(raw_dir: Path) -> list[Row]:
    """Собрать параметры кривой бескупонной доходности ОФЗ.

    В выгрузке ISS кривая пересчитывается в течение дня много раз;
    берётся последний снимок каждой даты — дневной срез, а не поминутный
    поток (внутридневные точки осознанно выброшены, они дают мегабайты
    на дату и ядру не нужны).

    Args:
        raw_dir: каталог оригиналов.

    Returns:
        Строки ряда: параметры `b1..b3`, `t1`, `g1..g9` на дату.
    """
    rows: list[Row] = []
    for path in sorted((raw_dir / "moex_zcyc").glob("zcyc_*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        block = payload["params"]
        columns = block["columns"]
        data = block["data"]
        if not data:
            logger.warning("пустой снимок %s", path.name)
            continue
        last = data[-1]
        record = dict(zip(columns, last))
        date = record["tradedate"]
        for name in columns[2:]:
            value = to_float(record.get(name))
            if value is not None:
                rows.append(Row(date, value, "nss_parameter", "RU", name))
    return rows


def read_mpl(raw_dir: Path) -> list[Row]:
    """Собрать макропруденциальные лимиты Банка России.

    Args:
        raw_dir: каталог оригиналов.

    Returns:
        Строки ряда: доля выдач, разрешённая регулятором, по виду кредита
        и диапазону ПДН, с датой начала действия.
    """
    path = raw_dir / "cbr_mpl" / "Macroprudential_limits_banks_01072025.xlsx"
    if not path.exists():
        logger.warning("нет файла %s", path)
        return []
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[Row] = []
    for sheet_name in book.sheetnames:
        for raw in book[sheet_name].iter_rows(values_only=True):
            if not raw or not str(raw[0] or "").strip().isdigit():
                continue
            start = raw[2]
            limit = to_float(raw[7])
            if limit is None or start is None:
                continue
            date = str(start)[:10]
            kind = re.sub(r"\s+", " ", str(raw[4] or "")).strip()[:80]
            pti = re.sub(r"\s+", " ", str(raw[5] or "")).strip()
            ltv = re.sub(r"\s+", " ", str(raw[6] or "")).strip()
            rows.append(Row(date, limit, "percent_of_new_loans", "RU",
                            f"{sheet_name}|ПДН {pti}|LTV {ltv}|{kind}"))
    return rows


def read_top10_deposit(raw_dir: Path) -> list[Row]:
    """Собрать максимальную ставку топ-10 банков по декадам.

    Args:
        raw_dir: каталог оригиналов.

    Returns:
        Строки ряда по декадам (первое число декады как дата).
    """
    path = raw_dir / "cbr_ref" / "avgprocstav_table.html"
    if not path.exists():
        logger.warning("нет файла %s", path)
        return []
    text = path.read_text(encoding="utf-8", errors="replace")
    pattern = re.compile(
        r"<td>(I{1,3})\.(\d{2})\.(\d{4})</td>\s*<td>([\d,]+)</td>")
    rows: list[Row] = []
    for decade, month, year, value in pattern.findall(text):
        number = to_float(html.unescape(value))
        if number is None:
            continue
        rows.append(Row(f"{year}-{month}-{DECADE[decade]}", number,
                        "percent_annual", "RU",
                        f"top10_max_deposit_rate|decade_{decade}"))
    rows.sort(key=lambda item: item.date)
    return rows


def read_gold(raw_dir: Path) -> list[Row]:
    """Собрать дневной ряд цены золота (LBMA PM fixing).

    Args:
        raw_dir: каталог оригиналов.

    Returns:
        Строки ряда в долларах за тройскую унцию.
    """
    path = raw_dir / "commodities" / "gold_pm.json"
    if not path.exists():
        logger.warning("нет файла %s", path)
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows: list[Row] = []
    for record in payload:
        values = record.get("v") or []
        if not values or values[0] is None:
            continue
        rows.append(Row(record["d"], float(values[0]), "usd_per_troy_ounce",
                        "WORLD", "gold_pm_fixing"))
    rows.sort(key=lambda item: item.date)
    return rows


def build(raw_dir: Path, out_dir: Path | None = None) -> None:
    """Собрать все справочные ряды.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения.
    """
    write_series("world_benchmarks.csv", read_world(raw_dir), out_dir)
    write_series("ofz_zcyc_params.csv", read_zcyc(raw_dir), out_dir)
    write_series("cbr_mpl_limits.csv", read_mpl(raw_dir), out_dir)
    write_series("cbr_deposit_rate_top10_decade.csv",
                 read_top10_deposit(raw_dir), out_dir)
    write_series("lbma_gold_pm_daily.csv", read_gold(raw_dir), out_dir)


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=raw_root())
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.raw_dir, args.out_dir)


if __name__ == "__main__":
    main()
