"""Ставки банков физическим лицам → `data/`.

Оригиналы (`~/raw-originals/finpilot-data/cbr/`), раздел статистики
`pdko/int_rat`:
  * `cbr_int_rat_loans_ind_new.xlsx` — ставки по кредитам физлицам,
    помесячно, по срокам (лист `ставки_руб`);
  * `cbr_int_rat_deposits.xlsx` — ставки по вкладам физлиц;
  * `cbr_int_rat_loans_ind_by_region_new.xlsx`,
    `cbr_int_rat_deposits_ind_by_region.xlsx` — то же по федеральным
    округам.

Берутся только рублёвые листы: валютные ставки ядру не нужны.

Запуск: `python -m tools.data.build_cbr_bank_rates`
"""
from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path

import openpyxl

from tools.data.common import Row, raw_root, to_float, write_series

logger = logging.getLogger(__name__)

MONTHS = {
    "январь": "01", "февраль": "02", "март": "03", "апрель": "04",
    "май": "05", "июнь": "06", "июль": "07", "август": "08",
    "сентябрь": "09", "октябрь": "10", "ноябрь": "11", "декабрь": "12",
}


def parse_month(label: str) -> str | None:
    """Перевести подпись «Январь 2014» в первый день месяца.

    Args:
        label: подпись строки или колонки.

    Returns:
        Дата ISO или None, если подпись не про месяц.
    """
    match = re.match(r"\s*([А-Яа-яё]+)\s+(\d{4})", str(label or ""))
    if not match:
        return None
    month = MONTHS.get(match.group(1).lower())
    if not month:
        return None
    return f"{match.group(2)}-{month}-01"


def clean(text: object) -> str:
    """Сжать подпись в одну строку без переносов и двойных пробелов.

    Args:
        text: исходная ячейка.

    Returns:
        Очищенная подпись.
    """
    return re.sub(r"\s+", " ", str(text or "").replace("\n", " ")).strip()


def column_groups(row: tuple) -> list[str]:
    """Растянуть объединённые подписи групп колонок вправо.

    В шапке ЦБ подпись группы («Физических лиц…», «Нефинансовых
    организаций…») стоит только над первой колонкой своего блока.

    Args:
        row: строка шапки с подписями групп.

    Returns:
        Подпись группы для каждой колонки строки.
    """
    groups: list[str] = []
    current = ""
    for cell in row:
        if clean(cell):
            current = clean(cell)
        groups.append(current)
    return groups


def read_national(path: Path, sheet: str,
                  group_prefix: str | None = None) -> list[Row]:
    """Прочитать помесячный лист «месяцы в строках, сроки в колонках».

    🔴 Лист вкладов ЦБ несёт ДВА блока колонок с одинаковыми подписями
    сроков: вклады физических лиц и депозиты нефинансовых организаций.
    Без фильтра по группе оба блока попадают в ряд под одинаковым
    `breakdown`, и ставка для бизнеса выдаётся за ставку для людей
    (поймано прогоном B2 17.09.2026).

    Args:
        path: путь к книге.
        sheet: имя рублёвого листа.
        group_prefix: начало подписи группы колонок, которую оставить;
            None — брать все колонки.

    Returns:
        Строки ряда с разрезом по сроку.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[sheet].iter_rows(values_only=True))
    groups = column_groups(rows[3])
    header = [clean(cell) for cell in rows[4]]
    out: list[Row] = []
    for raw in rows[5:]:
        date = parse_month(raw[0])
        if not date:
            continue
        for index in range(1, len(header)):
            if group_prefix and not (index < len(groups) and groups[index]
                                     .startswith(group_prefix)):
                continue
            term = header[index]
            value = to_float(raw[index]) if index < len(raw) else None
            if not term or value is None:
                continue
            out.append(Row(date, value, "percent_annual", "RU", term))
    return out


def read_regional(path: Path, sheet: str) -> list[Row]:
    """Прочитать лист «округа в строках, месяцы в колонках».

    Args:
        path: путь к книге.
        sheet: имя листа.

    Returns:
        Строки ряда с регионом и разрезом по сроку.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[sheet].iter_rows(values_only=True))
    header_index = next(i for i, raw in enumerate(rows)
                        if any(parse_month(cell) for cell in raw))
    header = rows[header_index]
    region_col = max(i for i, cell in enumerate(header)
                     if clean(cell).lower().startswith("федеральный округ"))
    term_col = region_col - 1
    dates = {i: parse_month(cell) for i, cell in enumerate(header)
             if parse_month(cell)}
    out: list[Row] = []
    term = ""
    for raw in rows[header_index + 1:]:
        if term_col < len(raw) and clean(raw[term_col]):
            term = clean(raw[term_col])
        region = clean(raw[region_col]) if region_col < len(raw) else ""
        if not region:
            continue
        for index, date in dates.items():
            value = to_float(raw[index]) if index < len(raw) else None
            if value is None:
                continue
            out.append(Row(date, value, "percent_annual", region, term))
    out.sort(key=lambda item: (item.date, item.region, item.breakdown))
    return out


def build(raw_dir: Path, out_dir: Path | None = None) -> None:
    """Собрать четыре ряда ставок.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения.
    """
    folder = raw_dir / "cbr"
    write_series(
        "cbr_loan_rate_individuals_month.csv",
        read_national(folder / "cbr_int_rat_loans_ind_new.xlsx", "ставки_руб"),
        out_dir,
    )
    write_series(
        "cbr_deposit_rate_individuals_month.csv",
        read_national(folder / "cbr_int_rat_deposits.xlsx", "ставки_руб",
                      group_prefix="Физических лиц"),
        out_dir,
    )
    write_series(
        "cbr_loan_rate_individuals_region_month.csv",
        read_regional(folder / "cbr_int_rat_loans_ind_by_region_new.xlsx",
                      "кредиты_физ_руб"),
        out_dir,
    )
    write_series(
        "cbr_deposit_rate_individuals_region_month.csv",
        read_regional(folder / "cbr_int_rat_deposits_ind_by_region.xlsx",
                      "вклады_физ_руб"),
        out_dir,
    )


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=raw_root())
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.raw_dir, args.out_dir)


if __name__ == "__main__":
    main()
