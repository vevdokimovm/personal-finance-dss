"""Инфляция, веса корзины, норма сбережения, прожиточный минимум → `data/`.

Оригиналы (`~/raw-originals/finpilot-data/rosstat/`):
  * `rosstat_ipc_mes_08-2026.xlsx` — ИПЦ помесячно, лист `01` (все товары);
  * `rosstat_ipc-KIPC_2010-2025.xlsx` — ИПЦ по классификатору КИПЦ;
  * `rosstat_Vesa-tov-KIPC_2012-2026.xlsx` — веса товаров в корзине;
  * `rosstat_urov_13kv_2kv2026.xlsx` — структура использования доходов
    поквартально (прирост финансовых активов = норма сбережения);
  * `rosstat_vpm_RF_kv.xlsx` и `rosstat_vpm-643_2021-2026.xlsx` —
    прожиточный минимум.

Из КИПЦ и весов берутся только верхние уровни классификатора (коды вида
`01` и `01.1`): нижние — тысячи товаров-представителей, для целей ядра
они не нужны и раздувают каталог.

Запуск: `python -m tools.data.build_rosstat_prices`
"""
from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path

import openpyxl

from tools.data.common import Row, raw_root, to_float, write_series

logger = logging.getLogger(__name__)

MONTHS = {
    "январь": "01", "февраль": "02", "март": "03", "апрель": "04",
    "май": "05", "июнь": "06", "июль": "07", "август": "08",
    "сентябрь": "09", "октябрь": "10", "ноябрь": "11", "декабрь": "12",
}
QUARTER_END = {"1": "03-31", "2": "06-30", "3": "09-30", "4": "12-31"}
MAX_CODE_DEPTH = 2


def build_cpi_month(source: Path, out_dir: Path | None) -> None:
    """ИПЦ помесячно, к предыдущему месяцу, РФ.

    Args:
        source: путь к `rosstat_ipc_mes_08-2026.xlsx`.
        out_dir: каталог назначения.
    """
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows = list(book["01"].iter_rows(values_only=True))
    years = [str(cell).strip() if cell else "" for cell in rows[3]]
    out: list[Row] = []
    for raw in rows[5:17]:
        label = str(raw[0] or "").strip().lower()
        month = MONTHS.get(label)
        if not month:
            continue
        for index, year in enumerate(years):
            if not re.fullmatch(r"\d{4}", year):
                continue
            value = to_float(raw[index]) if index < len(raw) else None
            if value is None:
                continue
            out.append(Row(f"{year}-{month}-01", value,
                           "index_prev_month_percent", "RU", "all_goods"))
    out.sort(key=lambda item: item.date)
    write_series("rosstat_cpi_month.csv", out, out_dir)


def _code_depth(code: str) -> int:
    """Глубина кода КИПЦ.

    Args:
        code: код вида `01.1.2`.

    Returns:
        Число уровней в коде.
    """
    return len([part for part in code.split(".") if part])


def header_year(cell: object) -> str:
    """Достать год из подписи колонки, снимая сноску Росстата.

    🔴 Колонка 2022 года подписана `20221)` — год плюс номер сноски.
    Строгое сравнение с четырьмя цифрами молча выбрасывало весь 2022 год
    (поймано прогоном B2 17.09.2026).

    Args:
        cell: ячейка строки заголовка.

    Returns:
        Год из четырёх цифр или пустая строка, если это не год.
    """
    match = re.fullmatch(r"(\d{4})(?:\d{1,2}\))?", str(cell or "").strip())
    return match.group(1) if match else ""


def build_cpi_kipc(source: Path, out_dir: Path | None) -> None:
    """ИПЦ по группам КИПЦ, годовой, верхние уровни классификатора.

    Args:
        source: путь к `rosstat_ipc-KIPC_2010-2025.xlsx`.
        out_dir: каталог назначения.
    """
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    out: list[Row] = []
    for sheet_name in book.sheetnames:
        if not re.fullmatch(r"\d{4}-\d{4}", sheet_name):
            continue
        rows = list(book[sheet_name].iter_rows(values_only=True))
        header_index = next(
            i for i, raw in enumerate(rows)
            if raw and str(raw[0] or "").startswith("Код товара")
        )
        years = [header_year(cell) for cell in rows[header_index]]
        for raw in rows[header_index + 1:]:
            code = str(raw[0] or "").strip()
            name = str(raw[1] or "").strip()
            if not name:
                continue
            if code and _code_depth(code) > MAX_CODE_DEPTH:
                continue
            label = f"{code} {name}".strip()
            for index, year in enumerate(years):
                if not year:
                    continue
                value = to_float(raw[index]) if index < len(raw) else None
                if value is None:
                    continue
                out.append(Row(f"{year}-12-31", value,
                               "index_dec_to_dec_percent", "RU", label))
    out.sort(key=lambda item: (item.date, item.breakdown))
    write_series("rosstat_cpi_kipc_year.csv", out, out_dir)


def weight_columns(sheet_name: str, header: tuple) -> dict[int, str]:
    """Разметить колонки листа весов корзины датами.

    Лист бывает двух видов: годовой (`2023`, одна колонка `СПР`) и
    сводный (`2012-2022 (до 01.04)`, по колонке на год). 🔴 Прогон B1
    читал у сводного листа только первую колонку и подписывал её 2012
    годом — веса 2013–2021 терялись (поймано прогоном B2 17.09.2026).
    Веса 2022 года «до 01.04» датируются 31 марта, чтобы не слиться
    с весами «с 01.04» того же года.

    Args:
        sheet_name: имя листа.
        header: строка заголовка таблицы.

    Returns:
        Отображение «номер колонки → дата ISO».
    """
    years = {index: header_year(cell) for index, cell in enumerate(header)
             if header_year(cell)}
    if not years:
        match = re.search(r"(\d{4})", sheet_name)
        return {2: f"{match.group(1)}-12-31"} if match else {}
    last = max(years.values())
    early = "до 01.04" in sheet_name
    return {index: f"{year}-03-31" if early and year == last else f"{year}-12-31"
            for index, year in years.items()}


def build_basket_weights(source: Path, out_dir: Path | None) -> None:
    """Веса групп товаров в потребительской корзине, по годам.

    Args:
        source: путь к `rosstat_Vesa-tov-KIPC_2012-2026.xlsx`.
        out_dir: каталог назначения.
    """
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    out: list[Row] = []
    for sheet_name in book.sheetnames:
        if sheet_name == "Содержание":
            continue
        rows = list(book[sheet_name].iter_rows(values_only=True))
        header_index = next(
            (i for i, raw in enumerate(rows)
             if raw and str(raw[0] or "").startswith("Код товара")), None
        )
        if header_index is None:
            continue
        columns = weight_columns(sheet_name, rows[header_index])
        for raw in rows[header_index + 1:]:
            code = str(raw[0] or "").strip()
            name = str(raw[1] or "").strip()
            if not name or (code and _code_depth(code) > MAX_CODE_DEPTH):
                continue
            for index, date in columns.items():
                value = to_float(raw[index]) if index < len(raw) else None
                if value is None:
                    continue
                out.append(Row(date, value, "percent_of_basket", "RU",
                               f"{code} {name}".strip()))
    out.sort(key=lambda item: (item.date, item.breakdown))
    write_series("rosstat_cpi_basket_weights_year.csv", out, out_dir)


def build_savings(source: Path, out_dir: Path | None) -> None:
    """Структура использования денежных доходов поквартально.

    Разрезы: `goods_and_services`, `obligatory_payments`,
    `financial_assets_change` (это и есть норма сбережения) и
    `cash_change`.

    Args:
        source: путь к `rosstat_urov_13kv_2kv2026.xlsx`.
        out_dir: каталог назначения.
    """
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows = list(book["Структура расходов"].iter_rows(values_only=True))
    names = ["income_total", "goods_and_services", "obligatory_payments",
             "financial_assets_change", "cash_change"]
    out: list[Row] = []
    year = ""
    for raw in rows:
        label = str(raw[0] or "").strip()
        if re.match(r"^\d{4}\s*год", label):
            year = label[:4]
            continue
        if not year:
            continue
        quarter = re.match(r"([1-4])\s*квартал", label)
        if quarter:
            date = f"{year}-{QUARTER_END[quarter.group(1)]}"
            suffix = ""
        elif label.lower().startswith("год"):
            date = f"{year}-12-31"
            suffix = "_annual"
        else:
            continue
        for offset, name in enumerate(names, start=1):
            value = to_float(raw[offset]) if offset < len(raw) else None
            if value is None:
                continue
            out.append(Row(date, value, "percent_of_income", "RU",
                           f"{name}{suffix}"))
    out.sort(key=lambda item: (item.date, item.breakdown))
    write_series("rosstat_income_use_quarter.csv", out, out_dir)


def build_subsistence(raw_dir: Path, out_dir: Path | None) -> None:
    """Прожиточный минимум: поквартально 2000–2020 и с 2021 по годам.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения.
    """
    groups = ["all", "working_age", "pensioners", "children"]
    out: list[Row] = []
    old = openpyxl.load_workbook(raw_dir / "rosstat" / "rosstat_vpm_RF_kv.xlsx",
                                 read_only=True, data_only=True)
    sheet = next(name for name in old.sheetnames if "Федерация" in name)
    year = ""
    for raw in old[sheet].iter_rows(values_only=True):
        label = str(raw[0] or "").strip()
        if re.fullmatch(r"\d{4}", label):
            year = label
            continue
        match = re.match(r"(I{1,3}V?|IV)\s*квартал", label)
        if not match or not year:
            continue
        quarter = {"I": "1", "II": "2", "III": "3", "IV": "4"}[match.group(1)]
        date = f"{year}-{QUARTER_END[quarter]}"
        for offset, group in enumerate(groups, start=1):
            value = to_float(raw[offset]) if offset < len(raw) else None
            if value is not None:
                out.append(Row(date, value, "rub_per_month", "RU", group))
    new = openpyxl.load_workbook(
        raw_dir / "rosstat" / "rosstat_vpm-643_2021-2026.xlsx",
        read_only=True, data_only=True)
    sheet = next(name for name in new.sheetnames if "Федерация" in name)
    for raw in new[sheet].iter_rows(values_only=True):
        label = str(raw[0] or "").strip()
        match = re.search(r"с 1 (\w+) (\d{4})", label)
        if not match:
            continue
        month = {"января": "01-01", "июня": "06-01"}.get(match.group(1), "01-01")
        date = f"{match.group(2)}-{month}"
        for offset, group in enumerate(groups, start=1):
            value = to_float(raw[offset]) if offset < len(raw) else None
            if value is not None:
                out.append(Row(date, value, "rub_per_month", "RU", group))
    out.sort(key=lambda item: (item.date, item.breakdown))
    write_series("rosstat_subsistence_minimum.csv", out, out_dir)


def build(raw_dir: Path, out_dir: Path | None = None) -> None:
    """Собрать все ряды цен и структуры доходов.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения.
    """
    folder = raw_dir / "rosstat"
    build_cpi_month(folder / "rosstat_ipc_mes_08-2026.xlsx", out_dir)
    build_cpi_kipc(folder / "rosstat_ipc-KIPC_2010-2025.xlsx", out_dir)
    build_basket_weights(folder / "rosstat_Vesa-tov-KIPC_2012-2026.xlsx", out_dir)
    build_savings(folder / "rosstat_urov_13kv_2kv2026.xlsx", out_dir)
    build_subsistence(raw_dir, out_dir)


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=raw_root())
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.raw_dir, args.out_dir)


if __name__ == "__main__":
    main()
