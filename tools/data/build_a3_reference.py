"""Ряды прогона A3 батча Г46 → `data/`.

  * `cbr_npf_returns_quarter.csv` — средневзвешенная доходность НПФ
    (пенсионные накопления и пенсионные резервы, до вознаграждения),
    активы и число участников ПДС — «Динамические ряды основных показателей
    деятельности НПФ» Банка России;
  * `rosstat_goal_prices_month.csv` — средние потребительские цены на
    позиции-цели (автомобили, обучение, детсад, отдых, КАСКО, аренда
    квартиры у частных лиц) по РФ,
    федеральным округам и субъектам, помесячно с 2022 года;
  * `cbr_fund_fees_class_snapshot.csv` — агрегаты комиссий ПИФ по типу
    фонда и классу активов из «Витрины данных ПИФ» Банка России: медиана
    и квартили вознаграждения УК и расходов, число фондов. Отдельные фонды
    в каталог не пишутся — нужен параметр класса, а не выбор фонда;
  * `un_wpp_life_table_rus_year.csv` — ожидаемая продолжительность
    предстоящей жизни `e(x)` и вероятность дожить до следующего возраста
    по полу и однолетнему возрасту, Россия, UN WPP 2024 (1950–2023).

Запуск: `python -m tools.data.build_a3_reference`
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import logging
import re
import statistics
from collections import defaultdict
from pathlib import Path

import openpyxl

from tools.data.common import Row, raw_root, to_float, write_series

logger = logging.getLogger(__name__)

GOAL_ITEMS = {
    7703: "car_domestic_new",
    7708: "car_foreign_new",
    7702: "car_imported_used",
    9882: "kasko_year",
    9933: "driving_course",
    9942: "university_state_semester",
    9941: "university_private_semester",
    9922: "college_semester",
    9921: "school_private_month",
    9513: "kindergarten_day",
    9931: "language_course_hour",
    9536: "trip_black_sea_russia",
    9541: "trip_turkey",
    9535: "trip_excursion_russia",
    9419: "rent_flat_1room_month",
    9421: "rent_flat_2room_month",
}
NPF_ROWS = {
    "Средневзвешенная доходность инвестирования пенсионных накоплений":
        ("percent_annual_gross", "pension_savings_ops_return"),
    "Средневзвешенная доходность от размещения пенсионных резервов":
        ("percent_annual_gross", "pension_reserves_npo_pds_return"),
    "Активы": ("rub_mln", "npf_assets"),
    "Количество участников (ПДС)": ("persons", "pds_participants"),
}
WPP_SEXES = ("Male", "Female")
WPP_KEEP_COLUMNS = ("ex", "px")


class NpfBuilder:
    """Ряды Банка России по негосударственным пенсионным фондам."""

    def __init__(self, raw_dir: Path) -> None:
        self.path = raw_dir / "npf" / "npf_stat.xlsx"

    def rows(self) -> list[Row]:
        """Прочитать выбранные строки листа динамических рядов.

        Returns:
            Точки рядов с датой конца квартала.
        """
        sheet = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        table = [list(r) for r in sheet.worksheets[0].iter_rows(values_only=True)]
        header = next(r for r in table if any(isinstance(c, dt.datetime) for c in r))
        start = next(i for i, c in enumerate(header) if isinstance(c, dt.datetime))
        result: list[Row] = []
        for record in table:
            name = next((c for c in record if isinstance(c, str)), "")
            target = self._match(name)
            if target is None:
                continue
            unit, label = target
            for date, cell in zip(header[start:], record[start:]):
                value = to_float(cell)
                if isinstance(date, dt.datetime) and value is not None:
                    result.append(Row(date.date().isoformat(), f"{value:.4f}", unit,
                                      "RU", f"cbr_npf|{label}"))
        return result

    @staticmethod
    def _match(name: str) -> tuple[str, str] | None:
        clean = name.strip()
        for prefix, target in NPF_ROWS.items():
            if prefix == "Активы" and clean != "Активы":
                continue
            if clean.startswith(prefix):
                return target
        return None


class GoalPriceBuilder:
    """Средние потребительские цены Росстата на позиции-цели."""

    SHEET = re.compile(r"^(\d{2})\((\d{4})\)$")

    def __init__(self, raw_dir: Path) -> None:
        self.folder = raw_dir / "goal_prices"

    def rows(self) -> list[Row]:
        """Собрать цены по всем месячным листам всех годовых файлов.

        Returns:
            Точки ряда: дата — первое число месяца, регион — название
            субъекта или федерального округа (`RU` для России).
        """
        result: dict[tuple[str, str, str], Row] = {}
        files = sorted(self.folder.glob("sred_potreb_cen_20[2-9][0-9].xlsx"))
        files += sorted(self.folder.glob("sred_potreb_cen_[0-9][0-9]-20[2-9][0-9].xlsx"))
        for path in files:
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for sheet in book.worksheets:
                match = self.SHEET.match(sheet.title)
                if not match:
                    continue
                date = f"{match.group(2)}-{match.group(1)}-01"
                for row in self._sheet_rows(sheet, date):
                    key = (row.date, row.region, row.breakdown)
                    if key in result:
                        logger.warning("дубль ключа %s в %s", key, path.name)
                    result[key] = row
        return [result[k] for k in sorted(result)]

    def _sheet_rows(self, sheet, date: str) -> list[Row]:
        iterator = sheet.iter_rows(values_only=True)
        codes: tuple | None = None
        found: list[Row] = []
        for record in iterator:
            if codes is None:
                if record and record[0] == "Код территории":
                    codes = record
                    next(iterator)  # строка названий позиций: коды достаточно
                continue
            if not record or not isinstance(record[0], (int, float)):
                continue
            region = self._region(int(record[0]), str(record[1]).strip())
            if region is None:
                continue
            for column, code in enumerate(codes):
                label = GOAL_ITEMS.get(code) if isinstance(code, int) else None
                value = to_float(record[column]) if label else None
                if value is not None:
                    found.append(Row(date, f"{value:.2f}", "rub", region, label))
        return found

    @staticmethod
    def _region(code: int, territory: str) -> str | None:
        """Оставить Россию, федеральные округа и субъекты, без городов.

        Коды территорий — ОКАТО: у субъекта (и автономного округа внутри области)
        восемь младших разрядов нулевые, у города нет; федеральные округа
        имеют короткие коды.
        """
        if code == 643:
            return "RU"
        if code < 1000 or code % 100_000_000 == 0:
            return territory
        return None


class FundFeeBuilder:
    """Агрегаты комиссий ПИФ из витрины данных Банка России."""

    CLASS_RULES = (
        ("money_market", r"денежн|ликвидн|сберегат|rusfar|ruonia|репо|накопит"),
        ("gold", r"золот|драгметал|драгоцен|gold"),
        ("bonds", r"облигац|rgbi|rucb|бонд|долгов"),
        ("equity", r"акци|imoex|mcftr|мосбирж|дивиденд|голуб"),
        ("mixed", r"смешан|сбалансир|всепогод|баланс"),
    )
    PERCENT = re.compile(r"(\d+(?:[.,]\d+)?)\s*%")

    def __init__(self, raw_dir: Path) -> None:
        self.path = raw_dir / "pif" / "mutual_fund_data.xlsx"

    def rows(self) -> list[Row]:
        """Посчитать медиану и квартили комиссий по группам.

        Returns:
            Точки на дату среза витрины; разрез — тип фонда, класс,
            показатель и статистика.
        """
        book = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        sheet = book.worksheets[0]
        table = list(sheet.iter_rows(values_only=True))
        date = self._snapshot_date(sheet.title, table)
        groups: dict[tuple[str, str], dict[str, list[float]]] = defaultdict(
            lambda: defaultdict(list))
        for record in table[5:]:
            if record[3] not in ("Биржевой", "Открытый") or record[5] != "Сформирован":
                continue
            fund_type = "etf_bpif" if record[3] == "Биржевой" else "open_opif"
            fund_class = self._classify(record)
            fees = {
                "management_fee": self._percent(record[13]),
                "infrastructure_fee": self._percent(record[15]),
                "max_other_expenses": self._percent(record[16]),
                "return_12m": to_float(record[24]),
            }
            for key in ((fund_type, fund_class), (fund_type, "all")):
                groups[key]["count"].append(1.0)
                for metric, value in fees.items():
                    if value is None:
                        continue
                    if metric == "return_12m":
                        value *= 100.0
                    groups[key][metric].append(value)
        return self._summarise(groups, date)

    def _summarise(self, groups, date: str) -> list[Row]:
        result: list[Row] = []
        for (fund_type, fund_class), metrics in sorted(groups.items()):
            base = f"cbr_pif_showcase|{fund_type}|{fund_class}"
            result.append(Row(date, float(len(metrics["count"])), "funds", "RU",
                              f"{base}|count"))
            for metric in ("management_fee", "infrastructure_fee",
                           "max_other_expenses", "return_12m"):
                values = sorted(metrics.get(metric, []))
                if len(values) < 3:
                    continue
                q1, q2, q3 = statistics.quantiles(values, n=4, method="inclusive")
                unit = "percent_12m" if metric == "return_12m" else "percent_of_nav_year"
                for stat, value in (("p25", q1), ("median", q2), ("p75", q3)):
                    result.append(Row(date, value, unit, "RU",
                                      f"{base}|{metric}|{stat}|n={len(values)}"))
        return result

    def _classify(self, record: tuple) -> str:
        text = f"{record[1]} {record[11]}".lower()
        for name, pattern in self.CLASS_RULES:
            if re.search(pattern, text):
                return name
        return "other"

    def _percent(self, cell: object) -> float | None:
        match = self.PERCENT.search(str(cell)) if cell is not None else None
        return float(match.group(1).replace(",", ".")) if match else None

    @staticmethod
    def _snapshot_date(title: str, table: list) -> str:
        for record in table[:3]:
            text = " ".join(str(c) for c in record if c)
            found = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", text)
            if found:
                return f"{found.group(3)}-{found.group(2)}-{found.group(1)}"
        found = re.search(r"(\d{2}) (\d{2}) (\d{4})", title)
        if not found:
            raise ValueError("не найдена дата среза витрины ПИФ")
        return f"{found.group(3)}-{found.group(2)}-{found.group(1)}"


class LifeTableBuilder:
    """Таблицы дожития UN WPP 2024 по России."""

    def __init__(self, raw_dir: Path) -> None:
        self.folder = raw_dir / "life_tables"

    def rows(self) -> list[Row]:
        """Прочитать отфильтрованные по России полные таблицы дожития.

        Returns:
            Точки `e(x)` и `p(x)` на 31 декабря года, разрез — пол и возраст.

        Raises:
            FileNotFoundError: если одного из полов нет в каталоге оригиналов.
        """
        result: list[Row] = []
        for sex in WPP_SEXES:
            path = self.folder / f"wpp2024_complete_{sex}_RUS_1950-2023.csv"
            with path.open(encoding="utf-8-sig", newline="") as handle:
                for record in csv.DictReader(handle):
                    if record["LocID"] != "643":
                        continue
                    age = int(record["AgeGrpStart"])
                    for column in WPP_KEEP_COLUMNS:
                        value = to_float(record[column])
                        if value is None:
                            continue
                        unit = "years" if column == "ex" else "probability"
                        result.append(Row(f"{record['Time']}-12-31", value, unit, "RU",
                                          f"un_wpp2024|{sex.lower()}|age={age:03d}|{column}"))
        return result


def build(raw_dir: Path, out_dir: Path | None) -> None:
    """Собрать все ряды прогона A3.

    Args:
        raw_dir: каталог оригиналов.
        out_dir: каталог назначения (по умолчанию `data/`).
    """
    write_series("cbr_npf_returns_quarter.csv", NpfBuilder(raw_dir).rows(), out_dir)
    write_series("rosstat_goal_prices_month.csv", GoalPriceBuilder(raw_dir).rows(), out_dir)
    write_series("cbr_fund_fees_class_snapshot.csv", FundFeeBuilder(raw_dir).rows(), out_dir)
    write_series("un_wpp_life_table_rus_year.csv", LifeTableBuilder(raw_dir).rows(), out_dir)


def main() -> None:
    """Точка входа командной строки."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=raw_root())
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    build(args.raw_dir, args.out_dir)


if __name__ == "__main__":
    main()
